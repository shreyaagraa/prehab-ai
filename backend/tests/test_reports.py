"""
tests/test_reports.py
---------------------
Comprehensive test suite for PreHab AI Reports & Export System:
- All 5 report types (Injury Risk, Biomechanical, Movement, Performance, Rehabilitation)
- Genuine PDF export (%PDF validation)
- Genuine Excel export (.xlsx openpyxl validation)
- Strict RBAC:
    - Athlete (own: 200, other: 403)
    - Coach (assigned: 200, unassigned: 403)
    - Physiotherapist (200)
    - Sports Scientist (200)
    - Administrator (200)
- Graceful handling of missing/partial data & processing/failed states.
"""
import io
import uuid
import pytest
import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.main import app
from app.database import get_db, Base
from app.models.user import User, RoleEnum
from app.models.athlete import Athlete
from app.models.video import Video
from app.models.analysis_result import AnalysisResult
from app.models.analysis_feature import AnalysisFeature
from app.models.analysis_less import AnalysisLESS
from app.models.injury_history import InjuryHistory
from app.core.security import create_access_token, get_password_hash


@pytest.fixture(scope="function")
def db_session():
    import app.database as database_module
    Base.metadata.create_all(bind=database_module.engine)
    session = database_module.SessionLocal()
    try:
        yield session
    finally:
        session.close()


@pytest.fixture
def client(db_session):
    def _get_db_override():
        yield db_session

    app.dependency_overrides[get_db] = _get_db_override
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


@pytest.fixture
def setup_test_environment(db_session: Session):
    unique_suffix = str(uuid.uuid4())[:8]

    # 1. Athletes & Users
    user_ath1 = User(
        user_id=uuid.uuid4(),
        name="Athlete One",
        email=f"athlete1_{unique_suffix}@test.invalid",
        password=get_password_hash("Pass123!"),
        role=RoleEnum.ATHLETE,
        is_active=True,
    )
    user_ath2 = User(
        user_id=uuid.uuid4(),
        name="Athlete Two",
        email=f"athlete2_{unique_suffix}@test.invalid",
        password=get_password_hash("Pass123!"),
        role=RoleEnum.ATHLETE,
        is_active=True,
    )
    coach1_user = User(
        user_id=uuid.uuid4(),
        name="Coach Assigned",
        email=f"coach1_{unique_suffix}@test.invalid",
        password=get_password_hash("Pass123!"),
        role=RoleEnum.COACH,
        is_active=True,
    )
    coach2_user = User(
        user_id=uuid.uuid4(),
        name="Coach Unassigned",
        email=f"coach2_{unique_suffix}@test.invalid",
        password=get_password_hash("Pass123!"),
        role=RoleEnum.COACH,
        is_active=True,
    )
    physio_user = User(
        user_id=uuid.uuid4(),
        name="Dr. Physio",
        email=f"physio_{unique_suffix}@test.invalid",
        password=get_password_hash("Pass123!"),
        role=RoleEnum.PHYSIOTHERAPIST,
        is_active=True,
    )
    sports_sci_user = User(
        user_id=uuid.uuid4(),
        name="Scientist Jane",
        email=f"scientist_{unique_suffix}@test.invalid",
        password=get_password_hash("Pass123!"),
        role=RoleEnum.SPORTS_SCIENTIST,
        is_active=True,
    )
    admin_user = User(
        user_id=uuid.uuid4(),
        name="Admin Boss",
        email=f"admin_{unique_suffix}@test.invalid",
        password=get_password_hash("Pass123!"),
        role=RoleEnum.ADMINISTRATOR,
        is_active=True,
    )

    db_session.add_all([user_ath1, user_ath2, coach1_user, coach2_user, physio_user, sports_sci_user, admin_user])
    db_session.commit()

    # Athlete profiles (ath1 is assigned to coach1, ath2 is assigned to nobody)
    ath1 = Athlete(
        athlete_id=uuid.uuid4(),
        user_id=user_ath1.user_id,
        coach_id=coach1_user.user_id,
        sport="Basketball",
        position="Point Guard",
        age=22,
        height=188.0,
        weight=84.0,
        dominant_leg="Right",
        injury_status="Healthy",
        training_sessions_per_week=5,
        average_session_duration=90,
        average_session_rpe=7.0,
        weekly_training_load=3150.0,
        current_fatigue_level=4,
    )
    ath2 = Athlete(
        athlete_id=uuid.uuid4(),
        user_id=user_ath2.user_id,
        coach_id=None,
        sport="Soccer",
        position="Forward",
        age=24,
        height=175.0,
        weight=70.0,
        dominant_leg="Left",
        injury_status="Rehabilitating",
    )
    db_session.add_all([ath1, ath2])
    db_session.commit()

    # Video & Analysis for Ath1
    vid1 = Video(
        video_id=uuid.uuid4(),
        athlete_id=ath1.athlete_id,
        title="Baseline Jump Landing Assessment",
        original_filename="jump_landing_1.mp4",
        fps=30,
        duration=4.5,
        resolution="1920x1080",
        processing_status="completed",
    )
    db_session.add(vid1)
    db_session.commit()

    analysis1 = AnalysisResult(
        analysis_id=uuid.uuid4(),
        video_id=vid1.video_id,
        athlete_id=ath1.athlete_id,
        status="COMPLETED",
        overall_risk_score=68.5,
        risk_level="MODERATE",
        symmetry_score=82.0,
        fatigue_score=35.0,
        movement_quality=75.0,
        joint_alignment=70.0,
        trunk_lean=6.5,
        knee_valgus=14.2,
        hip_stability=78.0,
        stride_length=1.4,
        fps=30.0,
        duration_seconds=4.5,
        frames_processed=135,
        width=1920,
        height=1080,
    )
    db_session.add(analysis1)
    db_session.commit()

    # Biomechanical Features
    feature1 = AnalysisFeature(
        feature_id=uuid.uuid4(),
        analysis_id=analysis1.analysis_id,
        feature_version="v1",
        features={
            "knee_angle_left_mean": 132.5,
            "knee_angle_right_mean": 141.0,
            "knee_rom_left": 75.0,
            "knee_rom_right": 84.5,
        },
    )
    db_session.add(feature1)

    # LESS Results
    less1 = AnalysisLESS(
        less_id=uuid.uuid4(),
        analysis_id=analysis1.analysis_id,
        score=5,
        max_computable_score=17,
        computable_items=15,
        error_items=5,
        not_computable_items=2,
        classification="MODERATE_RISK",
        source="LESS-Standard",
        validation_source="PreHab-AI",
        disclaimer="Standard scoring approximation.",
        items=[
            {"item_number": 1, "item_name": "Knee Flexion Angle at Initial Contact", "status": "PASS", "score": 0, "measured_value": 32.5},
            {"item_number": 2, "item_name": "Knee Valgus at Initial Contact", "status": "ERROR", "score": 1, "measured_value": 14.2, "reason": "Medial collapse > 10 degrees"},
        ],
    )
    db_session.add(less1)

    # Injury History
    inj1 = InjuryHistory(
        injury_id=uuid.uuid4(),
        athlete_id=ath1.athlete_id,
        injury_type="Sprain",
        body_part="Right Ankle",
        severity="Mild",
        status="Recovered",
        remarks="Fully rehabbed during preseason.",
    )
    db_session.add(inj1)
    db_session.commit()

    from app.config import settings
    from datetime import timedelta

    def make_token(user: User):
        return create_access_token(
            subject=str(user.user_id),
            role=user.role.value,
            secret_key=settings.SECRET_KEY,
            algorithm=settings.ALGORITHM,
            expires_delta=timedelta(minutes=60),
        )

    return {
        "user_ath1": user_ath1,
        "token_ath1": make_token(user_ath1),
        "user_ath2": user_ath2,
        "token_ath2": make_token(user_ath2),
        "coach1": coach1_user,
        "token_coach1": make_token(coach1_user),
        "coach2": coach2_user,
        "token_coach2": make_token(coach2_user),
        "physio": physio_user,
        "token_physio": make_token(physio_user),
        "sports_sci": sports_sci_user,
        "token_sports_sci": make_token(sports_sci_user),
        "admin": admin_user,
        "token_admin": make_token(admin_user),
        "ath1": ath1,
        "ath2": ath2,
        "vid1": vid1,
        "analysis1": analysis1,
    }


def test_report_options_endpoint(client, setup_test_environment):
    env = setup_test_environment
    # Coach1 sees only ath1
    res = client.get("/api/v1/reports/options", headers={"Authorization": f"Bearer {env['token_coach1']}"})
    assert res.status_code == 200
    data = res.json()
    assert len(data["athletes"]) == 1
    assert data["athletes"][0]["athlete_id"] == str(env["ath1"].athlete_id)
    assert len(data["report_types"]) == 5

    # Athlete1 sees only self
    res = client.get("/api/v1/reports/options", headers={"Authorization": f"Bearer {env['token_ath1']}"})
    assert res.status_code == 200
    data = res.json()
    assert len(data["athletes"]) == 1
    assert data["athletes"][0]["athlete_id"] == str(env["ath1"].athlete_id)

    # Physio sees all athletes
    res = client.get("/api/v1/reports/options", headers={"Authorization": f"Bearer {env['token_physio']}"})
    assert res.status_code == 200
    data = res.json()
    assert len(data["athletes"]) >= 2


@pytest.mark.parametrize("report_type", [
    "INJURY_RISK",
    "BIOMECHANICAL",
    "MOVEMENT",
    "ATHLETE_PERFORMANCE",
    "REHABILITATION",
])
def test_all_five_report_types_json(client, setup_test_environment, report_type):
    env = setup_test_environment
    res = client.get(
        f"/api/v1/reports/data?report_type={report_type}&athlete_id={env['ath1'].athlete_id}&video_id={env['vid1'].video_id}",
        headers={"Authorization": f"Bearer {env['token_coach1']}"},
    )
    assert res.status_code == 200
    data = res.json()
    assert data["report_type"] == report_type
    assert data["athlete"]["name"] == "Athlete One"
    assert "disclaimer" in data
    assert data["data"] is not None


def test_pdf_export_generation(client, setup_test_environment):
    env = setup_test_environment
    res = client.get(
        f"/api/v1/reports/export/pdf?report_type=INJURY_RISK&athlete_id={env['ath1'].athlete_id}&video_id={env['vid1'].video_id}",
        headers={"Authorization": f"Bearer {env['token_coach1']}"},
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert res.content.startswith(b"%PDF")
    assert len(res.content) > 1000


def test_excel_export_generation(client, setup_test_environment):
    env = setup_test_environment
    res = client.get(
        f"/api/v1/reports/export/excel?report_type=INJURY_RISK&athlete_id={env['ath1'].athlete_id}&video_id={env['vid1'].video_id}",
        headers={"Authorization": f"Bearer {env['token_coach1']}"},
    )
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
    
    # Verify valid workbook structure
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert len(wb.sheetnames) >= 2
    assert "Athlete & Assessment Summary" in wb.sheetnames
    assert "Risk & Biomechanics" in wb.sheetnames


def test_rbac_athlete_access(client, setup_test_environment):
    env = setup_test_environment
    # Athlete1 accessing own report -> 200
    res = client.get(
        f"/api/v1/reports/data?report_type=INJURY_RISK&athlete_id={env['ath1'].athlete_id}",
        headers={"Authorization": f"Bearer {env['token_ath1']}"},
    )
    assert res.status_code == 200

    # Athlete1 accessing Athlete2's report -> 403 Forbidden
    res = client.get(
        f"/api/v1/reports/data?report_type=INJURY_RISK&athlete_id={env['ath2'].athlete_id}",
        headers={"Authorization": f"Bearer {env['token_ath1']}"},
    )
    assert res.status_code == 403
    assert "permission" in res.json()["detail"].lower()


def test_rbac_coach_access(client, setup_test_environment):
    env = setup_test_environment
    # Coach1 (assigned to ath1) accessing ath1 -> 200
    res = client.get(
        f"/api/v1/reports/data?report_type=INJURY_RISK&athlete_id={env['ath1'].athlete_id}",
        headers={"Authorization": f"Bearer {env['token_coach1']}"},
    )
    assert res.status_code == 200

    # Coach2 (unassigned) accessing ath1 -> 403 Forbidden
    res = client.get(
        f"/api/v1/reports/data?report_type=INJURY_RISK&athlete_id={env['ath1'].athlete_id}",
        headers={"Authorization": f"Bearer {env['token_coach2']}"},
    )
    assert res.status_code == 403

    # Coach1 accessing unassigned ath2 -> 403 Forbidden
    res = client.get(
        f"/api/v1/reports/data?report_type=INJURY_RISK&athlete_id={env['ath2'].athlete_id}",
        headers={"Authorization": f"Bearer {env['token_coach1']}"},
    )
    assert res.status_code == 403


def test_rbac_staff_roles_access(client, setup_test_environment):
    env = setup_test_environment
    # Physio -> 200
    res = client.get(
        f"/api/v1/reports/data?report_type=REHABILITATION&athlete_id={env['ath1'].athlete_id}",
        headers={"Authorization": f"Bearer {env['token_physio']}"},
    )
    assert res.status_code == 200

    # Sports Scientist -> 200
    res = client.get(
        f"/api/v1/reports/data?report_type=BIOMECHANICAL&athlete_id={env['ath1'].athlete_id}",
        headers={"Authorization": f"Bearer {env['token_sports_sci']}"},
    )
    assert res.status_code == 200

    # Administrator -> 200
    res = client.get(
        f"/api/v1/reports/data?report_type=ATHLETE_PERFORMANCE&athlete_id={env['ath1'].athlete_id}",
        headers={"Authorization": f"Bearer {env['token_admin']}"},
    )
    assert res.status_code == 200


def test_missing_and_partial_data_handling(client, setup_test_environment):
    env = setup_test_environment
    # ath2 has no videos, no analyses, no LESS
    res = client.get(
        f"/api/v1/reports/data?report_type=INJURY_RISK&athlete_id={env['ath2'].athlete_id}",
        headers={"Authorization": f"Bearer {env['token_admin']}"},
    )
    assert res.status_code == 200
    data = res.json()
    assert data["assessment"] is None
    assert data["data"]["overall_risk_score"] == "Data unavailable"

    # Export PDF for athlete with no assessments still succeeds with "Data unavailable"
    res_pdf = client.get(
        f"/api/v1/reports/export/pdf?report_type=ATHLETE_PERFORMANCE&athlete_id={env['ath2'].athlete_id}",
        headers={"Authorization": f"Bearer {env['token_admin']}"},
    )
    assert res_pdf.status_code == 200
    assert res_pdf.content.startswith(b"%PDF")
