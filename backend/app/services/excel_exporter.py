"""
app/services/excel_exporter.py
------------------------------
Structured .xlsx spreadsheet generator using openpyxl.
Produces multi-sheet workbooks with styled headers, structured tables,
auto-fitted column widths, and genuine application data.
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.schemas.reports import FullReportPayload, ReportTypeEnum


# ── Styling constants ───────────────────────────────────────────────────────
HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
SECTION_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
CARD_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

WHITE_FONT_BOLD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
DARK_FONT_BOLD = Font(name="Calibri", size=11, bold=True, color="0F172A")
DARK_FONT_REGULAR = Font(name="Calibri", size=10, color="334155")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
DISCLAIMER_FONT = Font(name="Calibri", size=9, italic=True, color="64748B")

THIN_BORDER_SIDE = Side(border_style="thin", color="CBD5E1")
CELL_BORDER = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=THIN_BORDER_SIDE,
    bottom=THIN_BORDER_SIDE,
)


def _apply_table_header(row, titles: list[str]):
    for col_idx, text in enumerate(titles, start=1):
        cell = row[col_idx - 1]
        cell.value = text
        cell.fill = SUBHEADER_FILL
        cell.font = WHITE_FONT_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = CELL_BORDER


def _autofit_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len and len(val) < 80:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)


def generate_excel_report(payload: FullReportPayload) -> bytes:
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    data_dict = payload.data.model_dump() if hasattr(payload.data, "model_dump") else payload.data

    # =========================================================================
    # SHEET 1: Summary & Profile
    # =========================================================================
    ws_summary = wb.create_sheet(title="Athlete & Assessment Summary")
    ws_summary.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_summary.merge_cells("A1:E1")
    title_cell = ws_summary["A1"]
    title_cell.value = f"PREHAB AI — {payload.report_title.upper()}"
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws_summary.row_dimensions[1].height = 28

    # Subtitle
    ws_summary["A2"].value = f"Generated: {payload.generated_at.strftime('%Y-%m-%d %H:%M UTC')} | By: {payload.generated_by_name} ({payload.generated_by_role})"
    ws_summary["A2"].font = DISCLAIMER_FONT
    ws_summary.row_dimensions[2].height = 18

    # Athlete Info Section
    ws_summary["A4"].value = "ATHLETE INFORMATION"
    ws_summary["A4"].font = DARK_FONT_BOLD
    ws_summary["A4"].fill = SECTION_FILL

    ath = payload.athlete
    ath_rows = [
        ("Athlete Name", ath.name, "Sport", ath.sport),
        ("Email", ath.email or "Data unavailable", "Position", ath.position),
        ("Age (years)", ath.age, "Dominant Leg", ath.dominant_leg),
        ("Height (cm)", ath.height, "Weight (kg)", ath.weight),
        ("Injury Status", ath.injury_status, "Assigned Coach", ath.coach_name),
        ("Weekly Training Load", f"{ath.weekly_training_load} AU" if ath.weekly_training_load != "Data unavailable" else "Data unavailable", "Fatigue Level", f"{ath.current_fatigue_level} / 10" if ath.current_fatigue_level != "Data unavailable" else "Data unavailable"),
    ]

    for idx, (k1, v1, k2, v2) in enumerate(ath_rows, start=5):
        ws_summary[f"A{idx}"].value = k1
        ws_summary[f"A{idx}"].font = DARK_FONT_BOLD
        ws_summary[f"A{idx}"].border = CELL_BORDER

        ws_summary[f"B{idx}"].value = str(v1)
        ws_summary[f"B{idx}"].font = DARK_FONT_REGULAR
        ws_summary[f"B{idx}"].border = CELL_BORDER

        ws_summary[f"C{idx}"].value = k2
        ws_summary[f"C{idx}"].font = DARK_FONT_BOLD
        ws_summary[f"C{idx}"].border = CELL_BORDER

        ws_summary[f"D{idx}"].value = str(v2)
        ws_summary[f"D{idx}"].font = DARK_FONT_REGULAR
        ws_summary[f"D{idx}"].border = CELL_BORDER

    # Assessment Info Section (if available)
    ass = payload.assessment
    curr_row = 12
    if ass:
        ws_summary[f"A{curr_row}"].value = "ASSESSMENT METADATA"
        ws_summary[f"A{curr_row}"].font = DARK_FONT_BOLD
        ws_summary[f"A{curr_row}"].fill = SECTION_FILL
        curr_row += 1

        ass_rows = [
            ("Assessment Title", ass.title or ass.original_filename or "N/A", "Status", ass.analysis_status or "N/A"),
            ("Assessment Date", ass.assessment_date.strftime('%Y-%m-%d %H:%M UTC') if ass.assessment_date else "N/A", "Resolution", ass.resolution),
            ("Video FPS", ass.fps, "Duration (seconds)", ass.duration_seconds),
            ("Frames Processed", ass.frames_processed, "Video File", ass.original_filename or "N/A"),
        ]

        for k1, v1, k2, v2 in ass_rows:
            ws_summary[f"A{curr_row}"].value = k1
            ws_summary[f"A{curr_row}"].font = DARK_FONT_BOLD
            ws_summary[f"A{curr_row}"].border = CELL_BORDER

            ws_summary[f"B{curr_row}"].value = str(v1)
            ws_summary[f"B{curr_row}"].font = DARK_FONT_REGULAR
            ws_summary[f"B{curr_row}"].border = CELL_BORDER

            ws_summary[f"C{curr_row}"].value = k2
            ws_summary[f"C{curr_row}"].font = DARK_FONT_BOLD
            ws_summary[f"C{curr_row}"].border = CELL_BORDER

            ws_summary[f"D{curr_row}"].value = str(v2)
            ws_summary[f"D{curr_row}"].font = DARK_FONT_REGULAR
            ws_summary[f"D{curr_row}"].border = CELL_BORDER
            curr_row += 1

    _autofit_columns(ws_summary)

    # =========================================================================
    # SHEET 2: Risk & Biomechanics Findings
    # =========================================================================
    ws_risk = wb.create_sheet(title="Risk & Biomechanics")
    ws_risk.views.sheetView[0].showGridLines = True

    ws_risk["A1"].value = "RISK SCREENING & BIOMECHANICAL SCORES"
    ws_risk["A1"].font = TITLE_FONT

    # Overall Scores Card
    ws_risk["A3"].value = "Overall Risk Score:"
    ws_risk["A3"].font = DARK_FONT_BOLD
    ws_risk["B3"].value = str(data_dict.get("overall_risk_score", data_dict.get("latest_risk_score", "Data unavailable")))
    ws_risk["B3"].font = DARK_FONT_BOLD

    ws_risk["C3"].value = "Risk Classification:"
    ws_risk["C3"].font = DARK_FONT_BOLD
    ws_risk["D3"].value = str(data_dict.get("risk_level", data_dict.get("latest_risk_level", "Data unavailable")))
    ws_risk["D3"].font = DARK_FONT_BOLD

    ws_risk["A4"].value = "Symmetry Score:"
    ws_risk["A4"].font = DARK_FONT_BOLD
    ws_risk["B4"].value = str(data_dict.get("symmetry_score", "Data unavailable"))
    ws_risk["B4"].font = DARK_FONT_REGULAR

    ws_risk["C4"].value = "Fatigue Score:"
    ws_risk["C4"].font = DARK_FONT_BOLD
    ws_risk["D4"].value = str(data_dict.get("fatigue_score", "Data unavailable"))
    ws_risk["D4"].font = DARK_FONT_REGULAR

    for r in range(3, 5):
        for c in ["A", "B", "C", "D"]:
            ws_risk[f"{c}{r}"].border = CELL_BORDER

    # 5-Factor Risk Breakdown Table
    factors = data_dict.get("risk_factors", [])
    if factors:
        ws_risk["A6"].value = "CONTRIBUTING RISK FACTORS (5-FACTOR SCREENING MODEL)"
        ws_risk["A6"].font = DARK_FONT_BOLD
        ws_risk["A6"].fill = SECTION_FILL

        ws_risk.append(["Factor Key", "Factor Name", "Score (0-100)", "Status", "Factor Description"])
        _apply_table_header(ws_risk[ws_risk.max_row], ["Factor Key", "Factor Name", "Score (0-100)", "Status", "Factor Description"])

        for f in factors:
            ws_risk.append([
                f.get("factor_key", ""),
                f.get("factor_name", ""),
                str(f.get("contribution_score", "")),
                f.get("status", ""),
                f.get("description", ""),
            ])
            for cell in ws_risk[ws_risk.max_row]:
                cell.font = DARK_FONT_REGULAR
                cell.border = CELL_BORDER

    # Asymmetry Findings Table
    asyms = data_dict.get("asymmetries", [])
    if asyms:
        start_row = ws_risk.max_row + 2
        ws_risk[f"A{start_row}"].value = "BILATERAL ASYMMETRY FINDINGS"
        ws_risk[f"A{start_row}"].font = DARK_FONT_BOLD
        ws_risk[f"A{start_row}"].fill = SECTION_FILL

        ws_risk.append(["Body Region / Joint", "Left Side", "Right Side", "Difference", "Classification"])
        _apply_table_header(ws_risk[ws_risk.max_row], ["Body Region / Joint", "Left Side", "Right Side", "Difference", "Classification"])

        for a in asyms:
            ws_risk.append([
                a.get("body_region", ""),
                str(a.get("left_value", "")),
                str(a.get("right_value", "")),
                f"{a.get('asymmetry_percentage', '')}°",
                a.get("status", ""),
            ])
            for cell in ws_risk[ws_risk.max_row]:
                cell.font = DARK_FONT_REGULAR
                cell.border = CELL_BORDER

    _autofit_columns(ws_risk)

    # =========================================================================
    # SHEET 3: LESS Assessment (if applicable)
    # =========================================================================
    less_items = data_dict.get("less_items", [])
    if less_items:
        ws_less = wb.create_sheet(title="LESS Findings")
        ws_less.views.sheetView[0].showGridLines = True

        ws_less["A1"].value = "LANDING ERROR SCORING SYSTEM (LESS) ASSESSMENT"
        ws_less["A1"].font = TITLE_FONT

        less_score = data_dict.get("less_total_score", data_dict.get("less_score", "Data unavailable"))
        less_max = data_dict.get("less_max_score", data_dict.get("less_max", "Data unavailable"))
        less_class = data_dict.get("less_classification", "Data unavailable")

        ws_less["A3"].value = f"Total Score: {less_score} / {less_max}"
        ws_less["A3"].font = DARK_FONT_BOLD
        ws_less["C3"].value = f"Classification: {less_class}"
        ws_less["C3"].font = DARK_FONT_BOLD

        ws_less.append([])
        headers = ["Item #", "Item Name", "Status", "Score", "Measured Value", "Criterion Threshold", "Unit", "Reason / Description"]
        ws_less.append(headers)
        _apply_table_header(ws_less[ws_less.max_row], headers)

        for it in less_items:
            ws_less.append([
                it.get("item_number", ""),
                it.get("item_name", ""),
                it.get("status", ""),
                it.get("score") if it.get("score") is not None else "",
                str(it.get("measured_value", "")),
                str(it.get("criterion_threshold", "")),
                str(it.get("unit", "")),
                str(it.get("reason", "") or "Criteria met"),
            ])
            for cell in ws_less[ws_less.max_row]:
                cell.font = DARK_FONT_REGULAR
                cell.border = CELL_BORDER

        _autofit_columns(ws_less)

    # =========================================================================
    # SHEET 4: Recommendations
    # =========================================================================
    recs = data_dict.get("recommendations") or data_dict.get("corrective_recommendations")
    if recs and isinstance(recs, dict):
        ws_rec = wb.create_sheet(title="Corrective Recommendations")
        ws_rec.views.sheetView[0].showGridLines = True

        ws_rec["A1"].value = "RECOMMENDED CORRECTIVE ACTION PLAN"
        ws_rec["A1"].font = TITLE_FONT

        rec_items = (
            recs.get("exercise_recommendations", [])
            + recs.get("mobility_recommendations", [])
            + recs.get("strengthening_recommendations", [])
        )

        if rec_items:
            headers = ["Priority", "Category", "Title", "Why Recommended", "Description", "Exercises / Drills", "Frequency", "Duration", "Safety Note"]
            ws_rec.append([])
            ws_rec.append(headers)
            _apply_table_header(ws_rec[ws_rec.max_row], headers)

            for r in rec_items:
                ws_rec.append([
                    r.get("priority", "").upper(),
                    r.get("category", ""),
                    r.get("title", ""),
                    r.get("why", ""),
                    r.get("description", ""),
                    ", ".join(r.get("exercises", [])),
                    r.get("frequency", ""),
                    r.get("duration", ""),
                    r.get("safety_note", ""),
                ])
                for cell in ws_rec[ws_rec.max_row]:
                    cell.font = DARK_FONT_REGULAR
                    cell.border = CELL_BORDER

        _autofit_columns(ws_rec)

    # =========================================================================
    # SHEET 5: History & Longitudinal Tracking
    # =========================================================================
    hist = data_dict.get("assessment_history", [])
    injuries = data_dict.get("injury_records", [])

    if hist or injuries:
        ws_hist = wb.create_sheet(title="History & Progression")
        ws_hist.views.sheetView[0].showGridLines = True

        ws_hist["A1"].value = "LONGITUDINAL ASSESSMENT & INJURY HISTORY"
        ws_hist["A1"].font = TITLE_FONT

        if hist:
            ws_hist.append([])
            ws_hist["A3"].value = "ASSESSMENT HISTORY"
            ws_hist["A3"].font = DARK_FONT_BOLD
            ws_hist["A3"].fill = SECTION_FILL

            headers = ["Date", "Assessment Title", "Risk Score (0-100)", "Risk Level", "LESS Score", "Symmetry Score"]
            ws_hist.append(headers)
            _apply_table_header(ws_hist[ws_hist.max_row], headers)

            for h in hist:
                dt_str = datetime.fromisoformat(h.get("date")).strftime("%Y-%m-%d %H:%M") if h.get("date") else "N/A"
                ws_hist.append([
                    dt_str,
                    h.get("title", ""),
                    str(h.get("risk_score", "")),
                    str(h.get("risk_level", "")),
                    str(h.get("less_score", "")),
                    str(h.get("symmetry_score", "")),
                ])
                for cell in ws_hist[ws_hist.max_row]:
                    cell.font = DARK_FONT_REGULAR
                    cell.border = CELL_BORDER

        if injuries:
            start_row = ws_hist.max_row + 2
            ws_hist[f"A{start_row}"].value = "DOCUMENTED INJURY RECORDS"
            ws_hist[f"A{start_row}"].font = DARK_FONT_BOLD
            ws_hist[f"A{start_row}"].fill = SECTION_FILL

            headers = ["Injury Type", "Body Part", "Severity", "Injury Date", "Recovery Date", "Status", "Remarks"]
            ws_hist.append(headers)
            _apply_table_header(ws_hist[ws_hist.max_row], headers)

            for inj in injuries:
                ws_hist.append([
                    str(inj.get("injury_type", "")),
                    str(inj.get("body_part", "")),
                    str(inj.get("severity", "")),
                    str(inj.get("injury_date", "")),
                    str(inj.get("recovery_date", "")),
                    str(inj.get("status", "")),
                    str(inj.get("remarks", "")),
                ])
                for cell in ws_hist[ws_hist.max_row]:
                    cell.font = DARK_FONT_REGULAR
                    cell.border = CELL_BORDER

        _autofit_columns(ws_hist)

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
